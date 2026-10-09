import os
import sys
import re

# Ensure UTF-8 output encoding for console
if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import latex2mathml.converter
from lxml import etree

# ---------------------------------------------------------
# XSLT Transformer Initialization (LaTeX -> MathML -> OMML)
# ---------------------------------------------------------
XSL_CANDIDATES = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
]

_xslt_transform = None
for p in XSL_CANDIDATES:
    if os.path.exists(p):
        try:
            _xslt_doc = etree.parse(p)
            _xslt_transform = etree.XSLT(_xslt_doc)
            print(f"[Math Engine] Loaded XSLT from: {p}")
            break
        except Exception as e:
            print(f"[Math Engine] Error loading XSLT from {p}: {e}")

def latex_to_omml(latex_code):
    """Convert LaTeX formula string to Word OMML XML Element with n-ary operator operand fix."""
    if _xslt_transform is None:
        return None
    try:
        clean_code = latex_code.strip()
        # Common latex replacements for robustness
        clean_code = clean_code.replace(r'\longrightarrow', r'\rightarrow')
        clean_code = clean_code.replace(r'\longleftarrow', r'\leftarrow')
        clean_code = clean_code.replace(r'\mid', r'|')
        
        mathml_str = latex2mathml.converter.convert(clean_code)
        math_tree = etree.fromstring(mathml_str)
        omml = _xslt_transform(math_tree)
        
        root = omml.getroot()
        ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        
        # Fix empty <m:e/> inside <m:nary> (e.g. \sum, \prod, \int)
        # Prevents Word from rendering an empty dotted square placeholder ⬚
        nary_elements = root.xpath('.//m:nary', namespaces=ns)
        for nary in nary_elements:
            e_elem = nary.xpath('./m:e', namespaces=ns)
            if e_elem:
                e = e_elem[0]
                if len(e) == 0 and (e.text is None or not e.text.strip()):
                    parent = nary.getparent()
                    idx = parent.index(nary)
                    following_siblings = list(parent)[idx+1:]
                    for sib in following_siblings:
                        parent.remove(sib)
                        e.append(sib)

        omml_xml_str = etree.tostring(root, encoding='unicode')
        return parse_xml(omml_xml_str)
    except Exception as err:
        print(f"[Math Warning] Failed to convert LaTeX: '{latex_code}' | Error: {err}")
        return None

def set_cell_background(cell, fill_hex):
    """Set shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Configure Margins: Left 3cm, Right 2cm, Top 3cm, Bottom 3cm (Chuẩn đồ án HUIT)
    for section in doc.sections:
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(3.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # Configure Base Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0, 0, 0)
    
    # Configure complex script font
    rFonts = style.element.rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:cs'), 'Times New Roman')

    return doc

def add_formatted_runs(p, text, base_font_size=Pt(13), base_bold=False, base_italic=False, base_color=None):
    """Parse markdown bold (**text**), italic (*text*), code (`text`), and math ($text$)."""
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$[^\$]+?\$)')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        
        # Check for inline LaTeX math $...$
        if token.startswith('$') and token.endswith('$') and len(token) >= 2:
            math_content = token[1:-1].strip()
            omml_elem = latex_to_omml(math_content)
            if omml_elem is not None:
                p._element.append(omml_elem)
                continue
            
            # Fallback to plain text if math conversion failed
            run = p.add_run(math_content)
            run.font.name = 'Cambria Math'
            run.font.size = base_font_size
            run.italic = True
            continue

        is_bold = base_bold
        is_italic = base_italic
        is_code = False
        content = token

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            is_bold = True
            content = token[2:-2]
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            is_italic = True
            content = token[1:-1]
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            is_code = True
            content = token[1:-1]

        run = p.add_run(content)
        run.font.name = 'Times New Roman'
        run.font.size = base_font_size
        run.bold = is_bold
        run.italic = is_italic
        if is_code:
            run.font.name = 'Courier New'
        if base_color:
            run.font.color.rgb = base_color

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        add_formatted_runs(p, text, base_font_size=Pt(14), base_bold=True)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, text, base_font_size=Pt(13), base_bold=True)
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_runs(p, text, base_font_size=Pt(13), base_bold=True)
    elif level == 4:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        add_formatted_runs(p, text, base_font_size=Pt(13), base_bold=True, base_italic=False)
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        add_formatted_runs(p, text, base_font_size=Pt(13), base_bold=True)

def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(1.27)  # 1 tab indent
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    add_formatted_runs(p, text, base_font_size=Pt(13))

def add_math_block(doc, math_text):
    """Add a centered standalone block math equation (OMML)."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    
    omml_elem = latex_to_omml(math_text)
    if omml_elem is not None:
        p._element.append(omml_elem)
    else:
        run = p.add_run(math_text)
        run.font.name = 'Cambria Math'
        run.font.size = Pt(13)
        run.italic = True

def add_list_item(doc, text, is_bullet=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-0.63)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    
    bullet_symbol = "•  " if is_bullet else ""
    p.add_run(bullet_symbol)
    add_formatted_runs(p, text, base_font_size=Pt(13))

def add_image_with_caption(doc, img_rel_path, caption_text, base_dir):
    possible_paths = [
        os.path.join(base_dir, img_rel_path),
        os.path.join(base_dir, "docs", img_rel_path),
        os.path.join(base_dir, "docs", "image", img_rel_path),
        os.path.abspath(img_rel_path)
    ]
    
    found_path = None
    for p in possible_paths:
        if os.path.exists(p):
            found_path = p
            break
            
    if found_path:
        p_img = doc.add_paragraph()
        p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.first_line_indent = Cm(0)
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(found_path, width=Cm(15.5))
    else:
        print(f"Warning: Image not found: {img_rel_path}")

    # Check if caption is valid (not purely timestamp digits)
    if caption_text and not re.match(r'^\d{8,20}$', caption_text.strip()):
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.first_line_indent = Cm(0)
        p_cap.paragraph_format.line_spacing = 1.5
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        add_formatted_runs(p_cap, caption_text, base_font_size=Pt(12), base_italic=True)

def add_table_from_rows(doc, table_rows):
    if not table_rows:
        return
    
    num_rows = len(table_rows)
    num_cols = max(len(row) for row in table_rows)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    for r_idx, row in enumerate(table_rows):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            if is_header:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_background(cell, "F2F2F2")
                add_formatted_runs(p, val, base_font_size=Pt(12), base_bold=True)
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_runs(p, val, base_font_size=Pt(12))
            
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Add space after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(4)
    p_spacer.paragraph_format.space_after = Pt(6)
    p_spacer.paragraph_format.line_spacing = 1.5

def parse_and_append_markdown(doc, md_content, base_dir):
    lines = md_content.splitlines()
    in_math_block = False
    math_buffer = []
    in_table = False
    table_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Math block $$ ... $$
        if line.startswith('$$'):
            if in_math_block:
                math_buffer.append(line[2:].strip())
                math_text = " ".join(math_buffer).replace('$$', '').strip()
                if math_text:
                    add_math_block(doc, math_text)
                in_math_block = False
                math_buffer = []
            else:
                if line.endswith('$$') and len(line) > 2:
                    math_text = line[2:-2].strip()
                    add_math_block(doc, math_text)
                else:
                    in_math_block = True
                    math_buffer = [line[2:].strip()]
            i += 1
            continue
            
        if in_math_block:
            math_buffer.append(line)
            i += 1
            continue

        # Table processing
        if line.startswith('|') and line.endswith('|'):
            clean_cells = [c.strip() for c in line.split('|')[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in clean_cells if c):
                i += 1
                continue
            table_rows.append(clean_cells)
            in_table = True
            i += 1
            continue
        else:
            if in_table:
                add_table_from_rows(doc, table_rows)
                in_table = False
                table_rows = []

        # Blank line
        if not line:
            i += 1
            continue

        # Horizontal rule
        if line in ['---', '***', '___']:
            i += 1
            continue

        # Alert blocks (> [!NOTE], > ...)
        if line.startswith('>'):
            clean_line = line.lstrip('> ').strip()
            add_body_paragraph(doc, clean_line)
            i += 1
            continue

        # Headings
        if line.startswith('#'):
            h_match = re.match(r'^(#{1,6})\s+(.*)$', line)
            if h_match:
                level = len(h_match.group(1))
                h_text = h_match.group(2).strip()
                add_heading(doc, h_text, level)
                i += 1
                continue

        # Image ![caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if img_match:
            caption = img_match.group(1).strip()
            img_path = img_match.group(2).strip()
            
            next_caption = ""
            if i + 1 < len(lines):
                next_line = lines[i+1].strip()
                if (next_line.startswith('*Hình') or next_line.startswith('Hình') or next_line.startswith('*')) and next_line.endswith('*'):
                    next_caption = next_line.strip('* ')
                    i += 1
            final_caption = next_caption if next_caption else caption
            add_image_with_caption(doc, img_path, final_caption, base_dir)
            i += 1
            continue

        # Standalone caption
        if (line.startswith('*Hình') or line.startswith('Hình 1.')) and not line.startswith('!['):
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.first_line_indent = Cm(0)
            p_cap.paragraph_format.line_spacing = 1.5
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            add_formatted_runs(p_cap, line.strip('* '), base_font_size=Pt(12), base_italic=True)
            i += 1
            continue

        # List items
        if line.startswith('- ') or line.startswith('* '):
            list_text = line[2:].strip()
            add_list_item(doc, list_text, is_bullet=True)
            i += 1
            continue

        num_match = re.match(r'^\d+\.\s+(.*)$', line)
        if num_match:
            add_list_item(doc, line, is_bullet=False)
            i += 1
            continue

        # Regular body paragraph
        add_body_paragraph(doc, line)
        i += 1

    if in_table and table_rows:
        add_table_from_rows(doc, table_rows)

def build_merged_docx(file_list, output_docx_path, base_dir):
    doc = create_document()
    
    for idx, rel_file in enumerate(file_list):
        full_path = os.path.join(base_dir, rel_file)
        if not os.path.exists(full_path):
            print(f"File not found: {full_path}")
            continue
            
        print(f"Processing [{idx+1}/{len(file_list)}]: {rel_file}")
        with open(full_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
            
        parse_and_append_markdown(doc, md_content, base_dir)
        
        if idx < len(file_list) - 1:
            doc.add_page_break()

    try:
        doc.save(output_docx_path)
        file_size_kb = os.path.getsize(output_docx_path) / 1024
        print(f"Successfully generated docx (Size: {file_size_kb:.2f} KB): {output_docx_path}")
    except PermissionError:
        alt_path = output_docx_path.replace(".docx", "_FIXED.docx")
        print(f"[Warning] File is currently open in Microsoft Word: {output_docx_path}")
        print(f"[Fallback] Saving to alternative path: {alt_path}")
        doc.save(alt_path)
        file_size_kb = os.path.getsize(alt_path) / 1024
        print(f"Successfully generated docx (Size: {file_size_kb:.2f} KB): {alt_path}")
        output_docx_path = alt_path

    return output_docx_path

def verify_docx(docx_path):
    doc = Document(docx_path)
    print("=== DOCX VERIFICATION REPORT ===")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print(f"Total Sections: {len(doc.sections)}")
    omml_count = 0
    for p in doc.paragraphs:
        omml_count += len(p._element.xpath('.//m:oMath'))
    for t in doc.tables:
        omml_count += len(t._element.xpath('.//m:oMath'))
    print(f"Total OMML Math Elements Rendered: {omml_count}")
    for i, s in enumerate(doc.sections):
        print(f"Section {i+1} Margins: Left={s.left_margin.cm:.1f}cm, Right={s.right_margin.cm:.1f}cm, Top={s.top_margin.cm:.1f}cm, Bottom={s.bottom_margin.cm:.1f}cm")
    print("=== VERIFICATION COMPLETE ===")

FILE_LIST = [
    r"docs\abbreviation.md",
    r"docs\1.2.1.Relate_Work.md",
    r"docs\1.2.2.E2E_VLM_BCTC.md",
    r"docs\1.4_Range.md",
    r"docs\1.5_Technical.md",
    r"docs\Chapter_2_VLM_QLoRA.md",
    r"docs\Chapter_3_E2E_VLM_Hybrid_Routing.md",
    r"docs\Plan.md",
    r"docs\References.md"
]

if __name__ == "__main__":
    base_dir = r"c:\Users\Admin\HUIT - Học Tập\Năm 4\Computer_Vision"
    output_path = os.path.join(base_dir, "docs", "DE_CUONG_DO_AN_TONG_HOP.docx")
    final_path = build_merged_docx(FILE_LIST, output_path, base_dir)
    verify_docx(final_path)
