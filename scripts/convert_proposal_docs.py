import os
import sys
import re
import subprocess
import argparse
import zipfile
from lxml import etree

# Ensure UTF-8 output encoding for console
if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

WORKSPACE_DIR = r"c:\Users\Admin\HUIT - Học Tập\Năm 4\Computer_Vision"
PROPOSAL_DIR = os.path.join(WORKSPACE_DIR, "docs", "Proposal")

PROPOSAL_FILES = [
    "1.2.1.3_VLM_OCR_Free.md",
    "1.2.1.4_He_Thong_Parser_Va_Dinh_Tuyen_Lai.md",
    "1.2.2_Tinh_Moi_Neuro_Symbolic.md",
    "1.4_Doi_Tuong_Va_Pham_Vi.md",
    "1.5.2_Ky_Thuat_Nghien_Cuu.md",
    "2.1_Cau_Truc_Cac_Chuong.md",
    "2.2_Ke_Hoach_Thuc_Hien.md"
]

def ensure_font_rfonts(run, font_name="Times New Roman"):
    """Ensure font applies across ascii, hAnsi, eastAsia, and cs in OpenXML."""
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>')
        rPr.append(rFonts)
    else:
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rFonts.set(qn('w:eastAsia'), font_name)

def set_cell_background(cell, fill_hex):
    """Set shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def is_display_math_paragraph(p):
    """Check if paragraph represents a standalone display math equation."""
    # Check for <m:oMathPara> or sole <m:oMath> without significant text
    if len(p._element.xpath('.//m:oMathPara')) > 0:
        return True
    math_elems = p._element.xpath('.//m:oMath')
    if len(math_elems) > 0 and len(p.text.strip()) == 0:
        return True
    return False

def is_image_paragraph(p):
    """Check if paragraph contains an embedded image drawing."""
    return len(p._element.xpath('.//w:drawing | .//w:pict')) > 0

def is_caption_paragraph(p):
    """Check if paragraph is a caption for table or image."""
    text = p.text.strip()
    if p.style.name.lower().startswith('caption'):
        return True
    if re.match(r'^\*?(Hình|Bảng|Sơ đồ)\s+\d+[\.\d]*\s*[:\.]', text, re.IGNORECASE):
        return True
    return False

def post_process_docx(docx_path):
    """Apply strict academic formatting to a docx converted by Pandoc."""
    doc = Document(docx_path)

    # 1. Page Margins: Left 3cm, Right 2cm, Top 2cm, Bottom 2cm (Standard A4 academic formatting)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    # 2. Update styles
    for style_name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
        if style_name in doc.styles:
            st = doc.styles[style_name]
            st.font.name = 'Times New Roman'
            st.font.size = Pt(13)
            st.font.color.rgb = RGBColor(0, 0, 0)
            if hasattr(st, 'paragraph_format'):
                st.paragraph_format.line_spacing = 1.5

    # 3. Format Paragraphs
    for p in doc.paragraphs:
        style_name = p.style.name if p.style else 'Normal'
        text_strip = p.text.strip()

        # Check if heading
        if style_name.startswith('Heading'):
            level_match = re.search(r'\d+', style_name)
            level = int(level_match.group()) if level_match else 1
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            
            if level == 1:
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                h_size = Pt(14)
            elif level == 2:
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                h_size = Pt(13)
            elif level == 3:
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                h_size = Pt(13)
            else:
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                h_size = Pt(13)

            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = h_size
                r.bold = True
                ensure_font_rfonts(r, 'Times New Roman')
            continue

        # Check if display math block
        if is_display_math_paragraph(p):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                ensure_font_rfonts(r, 'Times New Roman')
            continue

        # Check if image paragraph
        if is_image_paragraph(p):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            continue

        # Check if caption paragraph
        if is_caption_paragraph(p):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.italic = True
                ensure_font_rfonts(r, 'Times New Roman')
            continue

        # Check if list item
        if style_name.startswith('List') or text_strip.startswith('•') or re.match(r'^\d+\.', text_strip):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.first_line_indent = Cm(-0.63)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
                ensure_font_rfonts(r, 'Times New Roman')
            continue

        # Check if IEEE Reference item: [1], [2], ...
        if re.match(r'^\[\d+\]', text_strip):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                ensure_font_rfonts(r, 'Times New Roman')
            continue

        # Check if Blockquote
        if style_name.startswith('Quote') or style_name.startswith('Block'):
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.right_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.italic = True
                ensure_font_rfonts(r, 'Times New Roman')
            continue

        # Standard body paragraph
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(1.27)  # 1 tab indent
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            ensure_font_rfonts(r, 'Times New Roman')

    # 4. Format Tables
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row in enumerate(table.rows):
            is_header = (r_idx == 0)
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if is_header:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for cell in row.cells:
                if is_header:
                    set_cell_background(cell, "F2F2F2")
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.2
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    if is_header:
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(12)
                        if is_header:
                            r.bold = True
                        ensure_font_rfonts(r, 'Times New Roman')

    # 5. Clean OMML math runs: strip m:sty val="p" so Word applies math typography & font
    for p in doc.paragraphs:
        for s in p._element.xpath('.//m:sty[@m:val="p"]'):
            s.getparent().remove(s)

    # Save back to docx_path
    try:
        doc.save(docx_path)
    except PermissionError:
        alt_path = docx_path.replace(".docx", "_FIXED.docx")
        print(f"[Warning] File locked: {docx_path}. Saving to: {alt_path}")
        doc.save(alt_path)
        docx_path = alt_path

    # Inject math font properties (<m:mathPr>) into word/settings.xml
    inject_math_properties(docx_path)

    # Build up equations into 2D Professional format via Word COM
    buildup_omaths_with_word(docx_path)

    return docx_path

def inject_math_properties(docx_path):
    """Ensure word/settings.xml specifies Cambria Math and 2D professional display properties."""
    math_pr_xml = """<m:mathPr xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
        <m:mathFont m:val="Cambria Math"/>
        <m:brkBin m:val="before"/>
        <m:brkBinSub m:val="--"/>
        <m:smallFrac m:val="0"/>
        <m:dispDef/>
        <m:lMargin m:val="0"/>
        <m:rMargin m:val="0"/>
        <m:defJc m:val="centerGroup"/>
        <m:wrapIndent m:val="1440"/>
        <m:intLim m:val="subSup"/>
        <m:naryLim m:val="undOvr"/>
    </m:mathPr>"""
    math_pr_elem = etree.fromstring(math_pr_xml)
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    
    temp_path = docx_path + ".tmp.zip"
    try:
        with zipfile.ZipFile(docx_path, "r") as zin:
            with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    content = zin.read(item.filename)
                    if item.filename == "word/settings.xml":
                        root = etree.fromstring(content)
                        m_prs = root.xpath('.//m:mathPr', namespaces=ns)
                        if not m_prs:
                            root.append(math_pr_elem)
                            content = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    zout.writestr(item, content)
        os.replace(temp_path, docx_path)
    except Exception as e:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def buildup_omaths_with_word(docx_path):
    """Use Microsoft Word COM to pre-compile equations into 2D Professional layout."""
    abs_path = os.path.abspath(docx_path)
    ps_cmd = f"""
    $w = New-Object -ComObject Word.Application
    $w.Visible = $false
    try {{
        $doc = $w.Documents.Open('{abs_path}')
        $doc.OMaths.BuildUp()
        $doc.Save()
        $doc.Close([ref]0)
    }} catch {{
    }} finally {{
        $w.Quit()
    }}
    """
    try:
        subprocess.run(["powershell", "-Command", ps_cmd], capture_output=True, timeout=15)
    except Exception:
        pass

def convert_md_to_docx(md_path, output_docx_path):
    """Run pandoc and post-process styling."""
    md_abs = os.path.abspath(md_path)
    out_abs = os.path.abspath(output_docx_path)
    workdir = os.path.dirname(md_abs)
    temp_docx = os.path.join(workdir, f"~temp_{os.path.splitext(os.path.basename(md_path))[0]}.docx")
    cmd = [
        "pandoc",
        md_abs,
        "-o", temp_docx,
        f"--resource-path={workdir}"
    ]
    
    result = subprocess.run(cmd, cwd=workdir, capture_output=True, text=True, encoding='utf-8')
    if result.returncode != 0:
        print(f"Pandoc error converting {md_path}: {result.stderr}")
        return None

    final_processed = post_process_docx(temp_docx)
    
    import shutil
    try:
        if os.path.exists(out_abs):
            os.remove(out_abs)
        os.replace(final_processed, out_abs)
        return out_abs
    except PermissionError:
        alt_path = out_abs.replace(".docx", "_FIXED.docx")
        print(f"[Warning] File locked by Word: {out_abs}. Saving to: {alt_path}")
        shutil.copy2(final_processed, alt_path)
        if os.path.exists(final_processed):
            try:
                os.remove(final_processed)
            except Exception:
                pass
        return alt_path

def verify_single_docx(docx_path):
    """Detailed verification of docx attributes against requirements."""
    doc = Document(docx_path)
    report = {}
    
    # Check Margins
    report['sections_count'] = len(doc.sections)
    sec = doc.sections[0]
    report['margins_cm'] = {
        'top': round(sec.top_margin.cm, 2),
        'bottom': round(sec.bottom_margin.cm, 2),
        'left': round(sec.left_margin.cm, 2),
        'right': round(sec.right_margin.cm, 2),
    }
    
    # Check Paragraphs
    report['paragraphs_count'] = len(doc.paragraphs)
    body_count = 0
    body_indented_count = 0
    body_justified_count = 0
    body_line_spacing_1_5_count = 0
    omml_count = 0
    headings_count = 0
    
    for p in doc.paragraphs:
        omml_count += len(p._element.xpath('.//m:oMath | .//m:oMathPara'))
        style_name = p.style.name if p.style else 'Normal'
        if style_name.startswith('Heading'):
            headings_count += 1
            continue
        if is_display_math_paragraph(p) or is_image_paragraph(p) or is_caption_paragraph(p):
            continue
        
        # Body text or list item
        body_count += 1
        indent_cm = p.paragraph_format.first_line_indent.cm if p.paragraph_format.first_line_indent else 0.0
        if round(indent_cm, 2) in [1.27, -0.63, -1.0]: # 1 tab or list hanging indent or reference hanging indent
            body_indented_count += 1
        if p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
            body_justified_count += 1
        if p.paragraph_format.line_spacing == 1.5:
            body_line_spacing_1_5_count += 1

    report['body_count'] = body_count
    report['body_indented_count'] = body_indented_count
    report['body_justified_count'] = body_justified_count
    report['body_line_spacing_1_5_count'] = body_line_spacing_1_5_count
    report['headings_count'] = headings_count
    report['omml_math_count'] = omml_count
    report['tables_count'] = len(doc.tables)

    # Check font in sample runs
    fonts_checked = set()
    sizes_checked = set()
    for p in doc.paragraphs[:20]:
        for r in p.runs:
            if r.font.name:
                fonts_checked.add(r.font.name)
            if r.font.size:
                sizes_checked.add(round(r.font.size.pt, 1))

    report['sample_fonts'] = list(fonts_checked)
    report['sample_sizes_pt'] = list(sizes_checked)
    report['file_size_kb'] = round(os.path.getsize(docx_path) / 1024, 2)
    return report

def merge_proposal_markdown(output_md_path):
    """
    Merge all Proposal markdown files in docs/Proposal into a single consolidated file,
    ordered by academic heading level sequence, eliminating duplicate top headings.
    """
    consolidated_parts = []
    
    for idx, filename in enumerate(PROPOSAL_FILES):
        file_path = os.path.join(PROPOSAL_DIR, filename)
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read().strip()
            
        if filename == "1.2.1.4_He_Thong_Parser_Va_Dinh_Tuyen_Lai.md":
            # Strip duplicate '# 1.2.' and '## 1.2.1.' headers as they already start 1.2.1.3
            lines = content.splitlines()
            cleaned_lines = []
            skip_header = True
            for line in lines:
                if skip_header and (line.strip().startswith("# 1.2.") or line.strip().startswith("## 1.2.1.")):
                    continue
                if skip_header and not line.strip():
                    continue
                skip_header = False
                cleaned_lines.append(line)
            content = "\n".join(cleaned_lines).strip()
            
        elif filename == "1.2.2_Tinh_Moi_Neuro_Symbolic.md":
            # Strip duplicate '# 1.2. TÌNH HÌNH NGHIÊN CỨU VÀ TÍNH MỚI CỦA ĐỀ TÀI'
            # as it already appears at the start of 1.2.1.3
            lines = content.splitlines()
            cleaned_lines = []
            skip_header = True
            for line in lines:
                if skip_header and line.strip().startswith("# 1.2."):
                    continue
                if skip_header and not line.strip():
                    continue
                skip_header = False
                cleaned_lines.append(line)
            content = "\n".join(cleaned_lines).strip()
            
        consolidated_parts.append(content)
        
    merged_content = "\n\n---\n\n".join(consolidated_parts) + "\n"
    
    # Sanitize display math blocks: remove internal blank lines that break Pandoc parser
    def sanitize_math_block(match):
        inner = match.group(1)
        clean_lines = [line.strip() for line in inner.splitlines() if line.strip()]
        return "$$\n" + "\n".join(clean_lines) + "\n$$"
    
    merged_content = re.sub(r'\$\$(.*?)\$\$', sanitize_math_block, merged_content, flags=re.DOTALL)
    
    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(merged_content)
        
    return output_md_path

def merge_and_convert(output_docx_name="Proposal_Tong_Hop.docx"):
    """Merge all markdown files and generate a single formatted docx."""
    merged_md_path = os.path.join(PROPOSAL_DIR, "PROPOSAL_CONSOLIDATED.md")
    output_docx_path = os.path.join(PROPOSAL_DIR, output_docx_name)
    
    print(f"=== 1. MERGING {len(PROPOSAL_FILES)} PROPOSAL MARKDOWN FILES ===")
    merge_proposal_markdown(merged_md_path)
    print(f"  [OK] Created consolidated Markdown: {merged_md_path}")
    
    print(f"\n=== 2. CONVERTING WITH PANDOC & POST-PROCESSING STYLING ===")
    final_docx = convert_md_to_docx(merged_md_path, output_docx_path)
    if final_docx and os.path.exists(final_docx):
        print(f"  [OK] Generated and styled: {final_docx}")
        report = verify_single_docx(final_docx)
        
        # Verify heading hierarchy
        doc = Document(final_docx)
        headings = []
        for p in doc.paragraphs:
            if p.style and p.style.name.startswith('Heading'):
                headings.append((p.style.name, p.text.strip()))
                
        print("\n=== 3. HEADING HIERARCHY VERIFICATION ===")
        print(f"Total Headings Detected: {len(headings)}")
        for style_name, h_text in headings:
            level_num = int(re.search(r'\d+', style_name).group()) if re.search(r'\d+', style_name) else 1
            prefix = "  " * (level_num - 1)
            print(f"{prefix}[{style_name}] {h_text}")
            
        print("\n=== 4. ACADEMIC FORMATTING EVIDENCE ===")
        print(f"  - File Name: {os.path.basename(final_docx)}")
        print(f"  - File Size: {report['file_size_kb']} KB")
        print(f"  - Total Paragraphs: {report['paragraphs_count']}")
        print(f"  - Total Tables: {report['tables_count']}")
        print(f"  - Math OMML Equations: {report['omml_math_count']}")
        print(f"  - Margins: Left={report['margins_cm']['left']}cm, Right={report['margins_cm']['right']}cm, Top={report['margins_cm']['top']}cm, Bottom={report['margins_cm']['bottom']}cm")
        print(f"  - Fonts Checked: {report['sample_fonts']}")
        print(f"  - Sizes Checked: {report['sample_sizes_pt']}")
        print(f"  - Body Paragraphs: {report['body_count']} (Justified: {report['body_justified_count']}, Indented 1.27cm: {report['body_indented_count']}, 1.5 Line Spacing: {report['body_line_spacing_1_5_count']})")
        return final_docx, report
    else:
        print("  [ERROR] Conversion failed.")
        return None, None

def convert_consolidated_directly(output_docx_name="Proposal_Tong_Hop.docx"):
    """Convert PROPOSAL_CONSOLIDATED.md directly to docx without re-merging sub-files."""
    merged_md_path = os.path.join(PROPOSAL_DIR, "PROPOSAL_CONSOLIDATED.md")
    output_docx_path = os.path.join(PROPOSAL_DIR, output_docx_name)
    
    print(f"=== 1. CONVERTING CONDENSED PROPOSAL_CONSOLIDATED.md WITH PANDOC & POST-PROCESSING ===")
    final_docx = convert_md_to_docx(merged_md_path, output_docx_path)
    if final_docx and os.path.exists(final_docx):
        print(f"  [OK] Generated and styled: {final_docx}")
        report = verify_single_docx(final_docx)
        
        # Verify heading hierarchy
        doc = Document(final_docx)
        headings = []
        for p in doc.paragraphs:
            if p.style and p.style.name.startswith('Heading'):
                headings.append((p.style.name, p.text.strip()))
                
        print("\n=== 2. HEADING HIERARCHY VERIFICATION ===")
        print(f"Total Headings Detected: {len(headings)}")
        for style_name, h_text in headings:
            level_num = int(re.search(r'\d+', style_name).group()) if re.search(r'\d+', style_name) else 1
            prefix = "  " * (level_num - 1)
            print(f"{prefix}[{style_name}] {h_text}")
            
        print("\n=== 3. ACADEMIC FORMATTING EVIDENCE ===")
        print(f"  - File Name: {os.path.basename(final_docx)}")
        print(f"  - File Size: {report['file_size_kb']} KB")
        print(f"  - Total Paragraphs: {report['paragraphs_count']}")
        print(f"  - Total Tables: {report['tables_count']}")
        print(f"  - Math OMML Equations: {report['omml_math_count']}")
        print(f"  - Margins: Left={report['margins_cm']['left']}cm, Right={report['margins_cm']['right']}cm, Top={report['margins_cm']['top']}cm, Bottom={report['margins_cm']['bottom']}cm")
        print(f"  - Fonts Checked: {report['sample_fonts']}")
        print(f"  - Sizes Checked: {report['sample_sizes_pt']}")
        print(f"  - Body Paragraphs: {report['body_count']} (Justified: {report['body_justified_count']}, Indented 1.27cm: {report['body_indented_count']}, 1.5 Line Spacing: {report['body_line_spacing_1_5_count']})")
        return final_docx, report
    else:
        print("  [ERROR] Conversion failed.")
        return None, None

def main():
    parser = argparse.ArgumentParser(description="Convert Proposal Markdown files to standard docx format")
    parser.add_argument("--test", action="store_true", help="Run conversion on a single test file and verify")
    parser.add_argument("--all", action="store_true", help="Convert all markdown files in docs/Proposal")
    parser.add_argument("--verify", action="store_true", help="Verify all existing docx files in docs/Proposal")
    parser.add_argument("--merge", action="store_true", help="Merge all markdown files in docs/Proposal into a single formatted docx")
    parser.add_argument("--consolidated", action="store_true", help="Convert PROPOSAL_CONSOLIDATED.md directly without re-merging")
    parser.add_argument("--output", type=str, default="Proposal_Tong_Hop.docx", help="Output filename for merged docx")
    parser.add_argument("-i", "--input", type=str, help="Input markdown file path to convert")
    parser.add_argument("-o", "--out", type=str, help="Output docx file path")
    args = parser.parse_args()

    if args.input:
        out_path = args.out if args.out else args.input.replace(".md", ".docx")
        print(f"=== CONVERTING CUSTOM FILE: {args.input} -> {out_path} ===")
        final_docx = convert_md_to_docx(args.input, out_path)
        if final_docx and os.path.exists(final_docx):
            report = verify_single_docx(final_docx)
            print("\n=== VERIFICATION EVIDENCE ===")
            print(f"- File: {os.path.basename(final_docx)} ({report['file_size_kb']} KB)")
            print(f"- Margins: Left={report['margins_cm']['left']}cm, Right={report['margins_cm']['right']}cm, Top={report['margins_cm']['top']}cm, Bottom={report['margins_cm']['bottom']}cm")
            print(f"- Total Paragraphs: {report['paragraphs_count']} (Body: {report['body_count']})")
            print(f"- Tables: {report['tables_count']}")
            print(f"- Math OMML: {report['omml_math_count']}")
            print(f"- Fonts: {report['sample_fonts']} | Sizes: {report['sample_sizes_pt']}")
        else:
            print("[ERROR] Conversion failed.")
        return

    if args.consolidated:
        convert_consolidated_directly(args.output)
        return

    if args.merge:
        merge_and_convert(args.output)
        return

    if args.test:
        test_file = "1.2.1.3_VLM_OCR_Free.md"
        md_path = os.path.join(PROPOSAL_DIR, test_file)
        out_path = os.path.join(PROPOSAL_DIR, test_file.replace(".md", ".docx"))
        print(f"=== RUNNING TEST CONVERSION: {test_file} ===")
        final_docx = convert_md_to_docx(md_path, out_path)
        if final_docx and os.path.exists(final_docx):
            print(f"[SUCCESS] Generated: {final_docx}")
            report = verify_single_docx(final_docx)
            print("\n=== VERIFICATION EVIDENCE ===")
            print(f"- Margins: Left={report['margins_cm']['left']}cm, Right={report['margins_cm']['right']}cm, Top={report['margins_cm']['top']}cm, Bottom={report['margins_cm']['bottom']}cm")
            print(f"- Line spacing 1.5 count: {report['body_line_spacing_1_5_count']} / {report['body_count']} body paragraphs")
            print(f"- Justified alignment count: {report['body_justified_count']} / {report['body_count']} body paragraphs")
            print(f"- 1-tab indent count: {report['body_indented_count']} / {report['body_count']} body paragraphs")
            print(f"- Fonts checked: {report['sample_fonts']}")
            print(f"- Font sizes checked: {report['sample_sizes_pt']}")
            print(f"- OMML Math equations rendered: {report['omml_math_count']}")
            print(f"- Tables rendered: {report['tables_count']}")
            print(f"- File size: {report['file_size_kb']} KB")
        else:
            print("[FAILED] Conversion failed.")
        return

    if args.all:
        print(f"=== CONVERTING ALL {len(PROPOSAL_FILES)} FILES IN {PROPOSAL_DIR} ===")
        for idx, filename in enumerate(PROPOSAL_FILES, 1):
            md_path = os.path.join(PROPOSAL_DIR, filename)
            out_path = os.path.join(PROPOSAL_DIR, filename.replace(".md", ".docx"))
            print(f"[{idx}/{len(PROPOSAL_FILES)}] Converting: {filename} -> {os.path.basename(out_path)}")
            final_docx = convert_md_to_docx(md_path, out_path)
            if final_docx and os.path.exists(final_docx):
                report = verify_single_docx(final_docx)
                print(f"  ✓ OK ({report['file_size_kb']} KB) | Math: {report['omml_math_count']} | Tables: {report['tables_count']} | Body: {report['body_count']}")
            else:
                print(f"  ✗ FAILED: {filename}")
        print("=== CONVERSION COMPLETE ===")
        return

    if args.verify:
        print(f"=== VERIFYING DOCX FILES IN {PROPOSAL_DIR} ===")
        for filename in PROPOSAL_FILES:
            docx_path = os.path.join(PROPOSAL_DIR, filename.replace(".md", ".docx"))
            if not os.path.exists(docx_path):
                print(f"[MISSING] {docx_path}")
                continue
            report = verify_single_docx(docx_path)
            print(f"\n[FILE] {filename.replace('.md', '.docx')}")
            print(f"  - Size: {report['file_size_kb']} KB")
            print(f"  - Margins: L={report['margins_cm']['left']}cm, R={report['margins_cm']['right']}cm, T={report['margins_cm']['top']}cm, B={report['margins_cm']['bottom']}cm")
            print(f"  - Fonts: {report['sample_fonts']} | Sizes: {report['sample_sizes_pt']}")
            print(f"  - Math OMML: {report['omml_math_count']} | Tables: {report['tables_count']}")
            print(f"  - Body: {report['body_count']} (Justified: {report['body_justified_count']}, Indented: {report['body_indented_count']}, 1.5 Line Spacing: {report['body_line_spacing_1_5_count']})")
        return

    parser.print_help()

if __name__ == "__main__":
    main()
